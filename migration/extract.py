#!/usr/bin/env python3
"""One-time migration: legacy dock workbook -> reservations, issues, import report, seed.sql.

Usage:
    python migration/extract.py --input data/dock_schedule.xlsx --out migration/out \
        [--generated-at 2026-09-22T00:00:00+00:00]

The annual sheets are a presentation grid, not data. A booking is a coloured bar (a merged
range and/or a run of same-coloured cells) on a berth row, with the vessel name somewhere in
it. This script turns bars into explicit reservations and logs every heuristic decision as an
import issue with a sheet!cell reference, so a person can review what was guessed.
Only dependency: openpyxl.
"""
from __future__ import annotations

import argparse
import calendar
import colorsys
import datetime as dt
import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

MONTHS = [m.upper() for m in calendar.month_name[1:]]
WEEKDAY_TOKENS = {"M", "T", "W", "TR", "TH", "F", "S", "SA", "SU"}
BERTH_RE = re.compile(r"^(?P<name>.+?)\s*-\s*(?P<len>\d+)\s*'\s*$")
POOLS = [  # shared resources: no length, no overlap rule
    ("NORTH FINGER PIERS", "north-finger-piers", "North Finger Piers"),
    ("SMALL CRAFT SLIPS", "small-craft-slips", "Small craft slips (institution boats)"),
]
UNASSIGNED = {"id": "unassigned", "name": "Unassigned (legacy rows)", "length_ft": None, "is_exclusive": 0}
VESSEL_RE = re.compile(r"^(R/V|M/V|S/V|S/Y|M/Y|OSV|OS/V|F/V|TUG|BARGE)\s+(.+)$", re.I)
PREFIX_DISPLAY = {"TUG": "Tug", "BARGE": "Barge"}
ANNOTATION_RE = re.compile(r"^(ETA|ETD|ARRIVAL|ARRIVES|DEPARTURE|DEPARTS|DELAYED|TOUCH AND GO|PROVISIONING|"
                           r"FUELING|FUEL TRUCK|BUNKERING|LOAD EQUIPMENT)\b", re.I)
CLOSURE_RE = re.compile(r"REPAIR|MAINTENANCE|REBUILD|NO DOCKING|NO USAGE|INSPECTION|\bTEST\b|REPLACEMENT|"
                        r"UTILITY WORK|PAVING|CONCRETE", re.I)
BACKGROUND_RGB = {"FFFFFF", "F2F2F2", "D9D9D9"}  # white and the grey shading used on quiet rows
THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
               "accent5", "accent6", "hlink", "folHlink"]  # Excel theme index -> scheme slot
DRAWING_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


# ----------------------------------------------------------------------------- helpers
def norm_text(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def norm_key(value) -> str:
    return norm_text(value).upper()


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower().replace("/", "")).strip("-")


def short_hash(*parts) -> str:
    return hashlib.sha1("|".join(str(p) for p in parts).encode()).hexdigest()[:12]


def is_text(value) -> bool:
    return isinstance(value, str) and value.strip() != "" and not value.startswith("=")


def vessel_display(key: str) -> str:
    m = VESSEL_RE.match(key)
    if not m:
        return key.title()
    prefix = m.group(1).upper()
    return f"{PREFIX_DISPLAY.get(prefix, prefix)} {m.group(2).title()}"


def classify(label_key: str) -> str:
    if VESSEL_RE.match(label_key):
        return "vessel"
    return "closure" if CLOSURE_RE.search(label_key) else "event"


def day_number(value) -> int | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


# ----------------------------------------------------------------------------- colours
def theme_palette(wb) -> list[str]:
    if not wb.loaded_theme:
        return []
    scheme = ET.fromstring(wb.loaded_theme).find(".//a:clrScheme", DRAWING_NS)
    palette = []
    for name in THEME_ORDER:
        el = scheme.find(f"a:{name}", DRAWING_NS) if scheme is not None else None
        child = el[0] if el is not None and len(el) else None
        palette.append(((child.get("lastClr") or child.get("val")) if child is not None else "000000").upper())
    return palette


def apply_tint(rgb: str, tint: float) -> str:
    if not tint:
        return rgb
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, lum, s = colorsys.rgb_to_hls(r, g, b)
    lum = lum * (1 + tint) if tint < 0 else lum * (1 - tint) + tint
    return "".join(f"{round(x * 255):02X}" for x in colorsys.hls_to_rgb(h, min(1.0, max(0.0, lum)), s))


def fill_rgb(cell, palette: list[str]) -> str | None:
    """Resolved solid fill as RRGGBB; None for no fill or background shading."""
    fill = getattr(cell, "fill", None)
    if fill is None or fill.fill_type != "solid" or fill.fgColor is None:
        return None
    c, rgb = fill.fgColor, None
    if c.type == "rgb" and isinstance(c.rgb, str):
        rgb = c.rgb[-6:].upper()
    elif c.type == "theme" and c.theme is not None and c.theme < len(palette):
        rgb = apply_tint(palette[c.theme], c.tint or 0.0)
    elif c.type == "indexed" and c.indexed is not None and c.indexed < 64:
        rgb = COLOR_INDEX[c.indexed][-6:].upper()
    return None if rgb is None or rgb in BACKGROUND_RGB else rgb


# ----------------------------------------------------------------------------- extraction
@dataclass
class Segment:
    resource_id: str
    lane: int
    start: dt.date
    end: dt.date
    label_key: str | None
    raw_label: str | None
    fill: str | None
    sources: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    merge_id: str | None = None


class Extractor:
    def __init__(self, path: Path):
        self.wb = load_workbook(path)  # formula mode on purpose: 1997-2001 day numbers have no cached values
        self.palette = theme_palette(self.wb)
        self.issues: list[dict] = []
        self.resources: dict[str, dict] = {}
        self.segments: list[Segment] = []
        self.carryover: dict[tuple[int, int], list[Segment]] = defaultdict(list)
        self.by_month: dict[tuple[int, int], list[Segment]] = defaultdict(list)
        self.stats = Counter()

    def issue(self, code, severity, message, sheet=None, cell=None, reservation_id=None):
        self.issues.append({"code": code, "severity": severity, "sheet": sheet, "cell": cell,
                            "message": message, "reservation_id": reservation_id})

    # -- resources ------------------------------------------------------------------
    @staticmethod
    def is_resource_label(label) -> bool:
        if not is_text(label):
            return False
        text = norm_text(label)
        return bool(BERTH_RE.match(text)) or any(text.upper().startswith(p) for p, _, _ in POOLS)

    def resource_for_label(self, label, sheet: str, row: int) -> str | None:
        if not self.is_resource_label(label):
            return None
        text = norm_text(label)
        m = BERTH_RE.match(text)
        if m:
            name, length = m.group("name").strip(), int(m.group("len"))
            rid = slug(name)
            known = self.resources.get(rid)
            if known is None:
                self.resources[rid] = {"id": rid, "name": name, "length_ft": length, "is_exclusive": 1,
                                       "sort_order": len(self.resources), "source": f"{sheet}!A{row}"}
            elif known["length_ft"] != length:
                self.issue("BERTH_LENGTH_CHANGED", "warn",
                           f"{name} labelled {length}' here but {known['length_ft']}' earlier", sheet, f"A{row}")
            return rid
        for prefix, rid, name in POOLS:
            if text.upper().startswith(prefix):
                self.resources.setdefault(rid, {"id": rid, "name": name, "length_ft": None, "is_exclusive": 0,
                                                "sort_order": len(self.resources), "source": f"{sheet}!A{row}"})
                return rid
        return None

    # -- sheets ---------------------------------------------------------------------
    def run(self):
        for sheet in sorted((s for s in self.wb.sheetnames if s.isdigit()), key=int):
            self.parse_sheet(sheet)
        self.compare_carryovers()
        for name in self.wb.sheetnames:
            if not name.isdigit():
                self.issue("SHEET_NOT_IMPORTED", "info", f"Sheet '{name}' is reference/reporting data", name)

    def parse_sheet(self, sheet: str):
        ws, year = self.wb[sheet], int(sheet)
        anchors = {}  # (row, col) -> (anchor_row, anchor_col, range_id, min_col, max_col)
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    anchors[(r, c)] = (rng.min_row, rng.min_col, f"{sheet}!{rng.coord}", rng.min_col, rng.max_col)
        headers = []
        for r in range(1, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if isinstance(a, str):
                month = next((i + 1 for i, m in enumerate(MONTHS) if norm_key(a).startswith(m)), None)
                if month:
                    headers.append((r, month, norm_text(a)))
        formula_days = 0
        for i, (hrow, month, header) in enumerate(headers):
            end_row = headers[i + 1][0] - 1 if i + 1 < len(headers) else ws.max_row
            header_year = int(m.group(1)) if (m := re.search(r"(\d{4})", header)) else None
            carry, block_year = False, year
            if header_year is not None and header_year != year:
                if month == 12 and header_year == year - 1 and i == 0:
                    carry, block_year = True, year - 1
                else:
                    self.issue("HEADER_YEAR_MISMATCH", "warn",
                               f"Header '{header}' in sheet {sheet}; used {year} from the sheet name", sheet, f"A{hrow}")
            rows = self.block_rows(ws, sheet, hrow, end_row, header)
            if not rows:
                continue
            day1_col, n_formulas = self.day_one(ws, sheet, hrow, rows[0][0], header, block_year, month)
            formula_days += n_formulas
            if day1_col is None:
                continue
            ndays = calendar.monthrange(block_year, month)[1]
            for row, rid, lane in rows:
                segs = self.parse_row(ws, sheet, row, rid, lane, block_year, month, day1_col, ndays, anchors)
                (self.carryover[(block_year, month)] if carry else self.segments).extend(segs)
                if not carry:
                    self.by_month[(block_year, month)].extend(segs)
            self.stats["carryover_blocks" if carry else "blocks"] += 1
        if formula_days:
            self.issue("DAY_NUMBERS_ARE_FORMULAS", "info",
                       f"{formula_days} day-number cells are uncached formulas; dates derived from column offsets",
                       sheet)

    def block_rows(self, ws, sheet, hrow, end_row, header) -> list[tuple[int, str, int]]:
        """Berth rows of one month block, plus unlabelled rows that carry bookings."""
        rows, lanes, last_rid = [], Counter(), None
        for r in range(hrow + 1, end_row + 1):
            label = ws.cell(r, 1).value
            rid = self.resource_for_label(label, sheet, r)
            texts = [norm_text(ws.cell(r, c).value) for c in range(2, ws.max_column + 1) if is_text(ws.cell(r, c).value)]
            if rid is not None:
                last_rid = rid
            elif last_rid is None or is_text(label):
                continue
            elif not any(not ANNOTATION_RE.match(t) for t in texts):
                if texts:
                    self.issue("NOTE_ROW", "info", f"Row of notes in '{header}' not tied to a berth: "
                               f"{'; '.join(texts)[:160]}", sheet, f"A{r}")
                continue
            elif not self.resources[last_rid]["is_exclusive"]:
                rid = last_rid
                self.issue("UNLABELLED_POOL_ROW", "info", f"Unlabelled row under {self.resources[rid]['name']} in "
                           f"'{header}'; imported as part of that shared resource", sheet, f"A{r}")
            else:
                rid = UNASSIGNED["id"]
                self.resources.setdefault(rid, {**UNASSIGNED, "sort_order": 999, "source": f"{sheet}!A{r}"})
                self.issue("UNLABELLED_ROW", "warn", f"Row with no berth label under "
                           f"{self.resources[last_rid]['name']} in '{header}' holds bookings "
                           f"({'; '.join(texts)[:120]}); imported to '{UNASSIGNED['name']}' for review. Possibly "
                           f"Marsh Landing, which appears in the 8YR summary but never as a row label.", sheet, f"A{r}")
            lane = lanes[rid]
            lanes[rid] += 1
            if lane > 0 and self.resources[rid]["is_exclusive"]:
                self.issue("DUPLICATE_BERTH_ROW", "warn", f"{self.resources[rid]['name']} has {lane + 1} rows in "
                           f"'{header}'; bookings on the extra row are imported as the same berth", sheet, f"A{r}")
            rows.append((r, rid, lane))
        return rows

    def day_one(self, ws, sheet, hrow, first_row, header, year, month) -> tuple[int | None, int]:
        """Column of day 1, by majority vote of every numeric day label around the header."""
        votes, labels, junk, formulas = Counter(), {}, 0, 0
        for r in range(max(1, hrow - 1), first_row):
            if r < hrow and self.is_resource_label(ws.cell(r, 1).value):
                continue
            for c in range(2, ws.max_column + 1):
                v = ws.cell(r, c).value
                n = day_number(v)
                if isinstance(v, str) and v.startswith("="):
                    formulas += 1
                elif n is not None and 1 <= n <= 31 and 2 <= c - n + 1 <= 12:
                    votes[c - n + 1] += 1
                    labels[c] = n
                elif is_text(v) and norm_key(v) not in WEEKDAY_TOKENS and r in (hrow, hrow + 1):
                    junk += 1
        if not votes:
            self.issue("NO_DAY_HEADER", "error", f"No day numbers found for '{header}'; block skipped", sheet, f"A{hrow}")
            return None, formulas
        day1 = votes.most_common(1)[0][0]
        if len(votes) > 1:
            detail = ", ".join(f"column {get_column_letter(k)} ({n} labels)" for k, n in votes.most_common())
            self.issue("HEADER_DAY_CONFLICT", "warn", f"Day numbers around '{header}' disagree on where day 1 is "
                       f"({detail}); used column {get_column_letter(day1)}", sheet, f"A{hrow}")
        if junk:
            self.issue("HEADER_ROW_TEXT", "warn", f"{junk} text cells sit in the day/weekday rows of '{header}'; "
                       f"ignored", sheet, f"A{hrow}")
        ndays = calendar.monthrange(year, month)[1]
        agreed = [n for c, n in labels.items() if c - n + 1 == day1]
        if invalid := sorted(n for n in agreed if n > ndays):
            self.issue("HEADER_INVALID_DAY", "warn", f"'{header}' lists day(s) {invalid} but the month has {ndays} "
                       f"days; cells under those columns are not imported", sheet, f"A{hrow}")
        if len(agreed) >= 20 and max(agreed) < ndays:
            self.issue("HEADER_MISSING_DAY", "warn", f"'{header}' stops at day {max(agreed)} of {ndays}; dates "
                       f"derived from column position", sheet, f"A{hrow}")
        return day1, formulas

    def parse_row(self, ws, sheet, row, rid, lane, year, month, day1_col, ndays, anchors) -> list[Segment]:
        segs: list[Segment] = []
        cur: Segment | None = None

        def close():
            nonlocal cur
            if cur is not None:
                segs.append(cur)
            cur = None

        for c in range(2, ws.max_column + 1):
            coord = f"{get_column_letter(c)}{row}"
            anchor = anchors.get((row, c))
            src = ws.cell(anchor[0], anchor[1]) if anchor else ws.cell(row, c)
            value, fill, merge_id = src.value, fill_rgb(src, self.palette), (anchor[2] if anchor else None)
            day = c - day1_col + 1
            if not 1 <= day <= ndays:
                own = ws.cell(row, c).value
                reaches_month = anchor is not None and anchor[3] <= day1_col + ndays - 1 and anchor[4] >= day1_col
                if is_text(own) and not reaches_month:
                    code, sev, where = ("PRE_GRID_LABEL", "info", "left of day 1") if day < 1 else \
                        ("OUT_OF_RANGE_TEXT", "warn", "past the last day")
                    self.issue(code, sev, f"'{norm_text(own)}' sits {where} of {calendar.month_name[month]} {year}; "
                               f"not imported as a booking", sheet, coord)
                continue
            date = dt.date(year, month, day)
            text = norm_text(value) if is_text(value) else None
            adjacent = cur is not None and cur.end == date - dt.timedelta(days=1)
            same_merge = adjacent and merge_id is not None and merge_id == cur.merge_id

            if text and ANNOTATION_RE.match(text):  # timing/service notes never create bookings
                note = f"{text} ({sheet}!{coord})"
                if adjacent and (fill is None or fill == cur.fill or same_merge):
                    cur.end = date
                    cur.notes.append(note)
                elif fill is not None:
                    close()
                    cur = Segment(rid, lane, date, date, None, None, fill, [f"{sheet}!{coord}"], [note], merge_id)
                else:
                    close()
                    self.issue("ORPHAN_NOTE", "info", f"Note '{text}' is not attached to any booking", sheet, coord)
                continue
            if text:
                key = norm_key(text)
                if same_merge or (adjacent and cur.label_key == key and
                                  (fill == cur.fill or fill is None or cur.fill is None)):
                    cur.end = date
                    cur.merge_id = merge_id or cur.merge_id
                elif adjacent and cur.label_key is None and cur.fill is not None and fill == cur.fill:
                    cur.label_key, cur.raw_label, cur.end = key, text, date  # the name sits mid-bar
                    cur.merge_id = merge_id or cur.merge_id
                    cur.sources.append(f"{sheet}!{coord}")
                else:
                    close()
                    cur = Segment(rid, lane, date, date, key, text, fill, [f"{sheet}!{coord}"], [], merge_id)
                continue
            if fill is not None:
                if adjacent and (fill == cur.fill or same_merge):
                    cur.end = date
                else:
                    close()
                    cur = Segment(rid, lane, date, date, None, None, fill, [f"{sheet}!{coord}"], [], merge_id)
                continue
            if same_merge:
                cur.end = date
                continue
            close()
        close()
        for s in segs:  # turn the first source into a cell range covering the bar within this month
            first = s.sources[0].split("!")[1]
            last = f"{get_column_letter(day1_col + s.end.day - 1)}{row}"
            if first != last:
                s.sources[0] = f"{sheet}!{first}:{last}"
        return segs

    # -- carry-over Decembers -------------------------------------------------------
    def compare_carryovers(self):
        key = lambda s: (s.resource_id, s.lane, s.start, s.end, s.label_key)
        for (year, month), copy in sorted(self.carryover.items()):
            original = {key(s) for s in self.by_month.get((year, month), [])}
            dup = {key(s) for s in copy}
            sheet = str(year + 1)
            if original == dup:
                self.issue("CARRYOVER_DUPLICATE", "info",
                           f"Sheet {sheet} repeats {calendar.month_name[month]} {year}; identical copy skipped", sheet)
                continue
            examples = "; ".join(f"{k[4] or 'unlabelled'} on {k[0]} {k[2]}..{k[3]}"
                                 for k in sorted(dup - original, key=lambda k: (k[2], k[0]))[:3])
            self.issue("CARRYOVER_MISMATCH", "warn",
                       f"Sheet {sheet} repeats {calendar.month_name[month]} {year} but disagrees with sheet {year} "
                       f"({len(dup - original)} bookings only in the copy, {len(original - dup)} only in the "
                       f"original; e.g. {examples}). Copy skipped; sheet {year} is the source of truth.", sheet)

    # -- stitching across month boundaries ------------------------------------------
    def stitched(self) -> list[Segment]:
        out: list[Segment] = []
        for s in sorted(self.segments, key=lambda s: (s.resource_id, s.lane, s.start)):
            prev = out[-1] if out else None
            if prev and (prev.resource_id, prev.lane) == (s.resource_id, s.lane) \
                    and s.start == prev.end + dt.timedelta(days=1):
                same_label = s.label_key is not None and s.label_key == prev.label_key
                colour_continues = s.label_key is None and s.fill is not None and s.fill == prev.fill and s.start.day == 1
                if same_label or colour_continues:
                    prev.end = s.end
                    prev.sources += s.sources
                    prev.notes += s.notes
                    continue
            out.append(s)
        return out

    # -- vessel reference sheets ----------------------------------------------------
    def reference_vessels(self) -> dict[str, dict]:
        found: dict[str, list[tuple[int, str]]] = defaultdict(list)
        for sheet in ("Science", "Yachts"):
            if sheet not in self.wb.sheetnames:
                continue
            ws = self.wb[sheet]
            for r in range(1, ws.max_row + 1):
                a = ws.cell(r, 1).value
                if not is_text(a):
                    continue
                text = norm_text(a)
                if text.upper().startswith("LOA:"):
                    self.issue("ORPHAN_SPEC_ROW", "info", f"'{text}' has no vessel name on its row; not attached",
                               sheet, f"A{r}")
                    continue
                m = re.match(r"^(.*?)\s+(\d+)\s*'$", text)
                if not m or not VESSEL_RE.match(m.group(1)):
                    continue
                key = norm_key(m.group(1))
                found[key].append((int(m.group(2)), f"{sheet}!A{r}"))
                for c in range(2, ws.max_column + 1):
                    v = ws.cell(r, c).value
                    loa = re.search(r"LOA:\s*(\d+)\s*'", v) if isinstance(v, str) else None
                    if loa and int(loa.group(1)) != int(m.group(2)):
                        found[key].append((int(loa.group(1)), f"{sheet}!{get_column_letter(c)}{r}"))
        out = {}
        for key, cands in found.items():
            lengths = sorted({n for n, _ in cands})
            listing = ", ".join(f"{n}' ({src})" for n, src in cands)
            disputed = len(lengths) > 1
            if disputed:
                sheet, cell = cands[0][1].split("!")
                self.issue("VESSEL_LENGTH_DISPUTED", "warn", f"{vessel_display(key)} has conflicting lengths: "
                           f"{listing}; the larger value is used for fit checks", sheet, cell)
            out[key] = {"length_ft": max(lengths), "status": "disputed" if disputed else "known",
                        "note": f"Disputed: {listing}" if disputed else None,
                        "source": ", ".join(src for _, src in cands)}
        return out

    # -- assemble -------------------------------------------------------------------
    def build(self, generated_at: str) -> dict:
        self.run()
        segments = self.stitched()
        ref = self.reference_vessels()
        vessels: dict[str, dict] = {}

        def vessel_for(key: str) -> dict:
            if key not in vessels:
                vid = slug(key)
                while any(v["id"] == vid for v in vessels.values()):
                    vid += "-x"
                info = ref.get(key)
                vessels[key] = {"id": vid, "name": vessel_display(key), "name_key": key,
                                "length_ft": info["length_ft"] if info else None,
                                "length_status": info["status"] if info else "unknown",
                                "length_note": info["note"] if info else None,
                                "source": info["source"] if info else "schedule only"}
            return vessels[key]

        for key in sorted(ref):
            vessel_for(key)
        reservations = []
        for s in segments:
            kind = "hold" if s.label_key is None else classify(s.label_key)
            vessel = vessel_for(s.label_key) if kind == "vessel" else None
            res = {"id": "lg-" + short_hash(s.resource_id, s.lane, s.start, s.end, s.label_key, s.sources[0]),
                   "berth_id": s.resource_id, "kind": kind, "vessel_id": vessel["id"] if vessel else None,
                   "title": None if vessel else ("Unlabelled block" if kind == "hold" else s.raw_label),
                   "start_date": s.start.isoformat(), "end_date": s.end.isoformat(),
                   "notes": "; ".join(s.notes) or None, "origin": "legacy",
                   "source_ref": "; ".join(s.sources), "lane": s.lane}
            reservations.append(res)
            if kind == "hold":
                sheet, cell = s.sources[0].split("!")
                self.issue("UNLABELLED_BLOCK", "info", f"Coloured cells with no name on "
                           f"{self.resources[s.resource_id]['name']} {res['start_date']}..{res['end_date']}; "
                           f"imported as a hold", sheet, cell, res["id"])
        reservations.sort(key=lambda r: (r["start_date"], r["berth_id"], r["id"]))
        issues = detect_issues(reservations, self.resources, {v["id"]: v for v in vessels.values()}, generated_at)
        unknown = sum(v["length_status"] == "unknown" for v in vessels.values())
        self.issue("VESSEL_LENGTH_UNKNOWN", "info", f"{unknown} vessels have no length in Science/Yachts; their "
                   f"bookings show 'fit unverified' until someone enters a length")
        self.issue("SUMMARY_NOT_IMPORTED", "info", "'8YR Dock Summary' credits North Pier West with 648 days in "
                   "2012, more days than a year has, which implies several vessels share that pier. The app treats "
                   "each berth as one booking at a time; confirm with the dock coordinator.", "8YR Dock Summary")
        span = lambda r: (dt.date.fromisoformat(r["end_date"]) - dt.date.fromisoformat(r["start_date"])).days + 1
        summary = {
            "generated_at": generated_at, "source_file": None,
            "sheets_parsed": sum(s.isdigit() for s in self.wb.sheetnames),
            "month_blocks": self.stats["blocks"], "carryover_blocks_skipped": self.stats["carryover_blocks"],
            "reservations": len(reservations),
            "reservations_by_kind": dict(Counter(r["kind"] for r in reservations)),
            "vessels": len(vessels),
            "vessels_by_length_status": dict(Counter(v["length_status"] for v in vessels.values())),
            "berths": len(self.resources),
            "schedule_issues_by_type": dict(Counter(i["type"] for i in issues)),
            "import_issues_by_code": dict(Counter(i["code"] for i in self.issues)),
            "data_range": {"from": reservations[0]["start_date"], "to": max(r["end_date"] for r in reservations)},
            "max_span_days": max(span(r) for r in reservations),
        }
        return {"berths": sorted(self.resources.values(), key=lambda b: b["sort_order"]),
                "vessels": sorted(vessels.values(), key=lambda v: v["name_key"]),
                "reservations": reservations, "issues": issues, "import_issues": self.issues, "summary": summary}


# ----------------------------------------------------------------------------- schedule rules
def shared_days(a: dict, b: dict) -> int:
    """Days two INCLUSIVE ranges have in common (0 when they don't overlap)."""
    start, end = max(a["start_date"], b["start_date"]), min(a["end_date"], b["end_date"])
    return 0 if start > end else (dt.date.fromisoformat(end) - dt.date.fromisoformat(start)).days + 1


def detect_issues(reservations: list[dict], berths: dict, vessels: dict, created_at: str) -> list[dict]:
    """Same rules the Worker enforces on writes (PRD section 6)."""
    out = []

    def add(kind, a, b, berth_id, details):
        out.append({"id": f"{kind[:2]}-" + short_hash(kind, a["id"], b["id"] if b else ""), "type": kind,
                    "reservation_id": a["id"], "other_reservation_id": b["id"] if b else None, "berth_id": berth_id,
                    "start_date": max(a["start_date"], b["start_date"]) if b else a["start_date"],
                    "end_date": min(a["end_date"], b["end_date"]) if b else a["end_date"],
                    "details": details, "created_at": created_at})

    by_berth, by_vessel = defaultdict(list), defaultdict(list)
    for r in reservations:
        by_berth[r["berth_id"]].append(r)
        if r["vessel_id"]:
            by_vessel[r["vessel_id"]].append(r)
    for berth_id, rows in by_berth.items():  # rule O: exclusive berths, any shared day conflicts
        if not berths[berth_id]["is_exclusive"]:
            continue
        active: list[dict] = []
        for r in sorted(rows, key=lambda r: (r["start_date"], r["id"])):
            active = [a for a in active if a["end_date"] >= r["start_date"]]
            for a in active:
                add("overlap", a, r, berth_id, {"shared_days": shared_days(a, r)})
            active.append(r)
    for r in reservations:  # rule F: vessel longer than the berth
        if r["kind"] == "vessel":
            v, b = vessels[r["vessel_id"]], berths[r["berth_id"]]
            if v["length_ft"] is not None and b["length_ft"] is not None and v["length_ft"] > b["length_ft"]:
                add("fit", r, None, r["berth_id"], {"vessel_length_ft": v["length_ft"],
                                                    "berth_length_ft": b["length_ft"],
                                                    "length_status": v["length_status"]})
    for rows in by_vessel.values():  # rule V: same vessel at two berths for 2+ days (1 day = shift day)
        rows = sorted(rows, key=lambda r: (r["start_date"], r["id"]))
        for i, a in enumerate(rows):
            for b in rows[i + 1:]:
                if b["start_date"] > a["end_date"]:
                    break
                if b["berth_id"] != a["berth_id"] and shared_days(a, b) >= 2:
                    add("vessel_double", a, b, None, {"berths": [a["berth_id"], b["berth_id"]],
                                                      "shared_days": shared_days(a, b)})
    return sorted(out, key=lambda i: (i["start_date"], i["type"], i["id"]))


# ----------------------------------------------------------------------------- SQL output
def sql_value(v) -> str:
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, (dict, list)):
        v = json.dumps(v, separators=(",", ":"))
    return "'" + str(v).replace("'", "''") + "'"


def inserts(table: str, cols: list[str], rows: list[dict], chunk: int = 100) -> list[str]:
    return [f"INSERT INTO {table} ({','.join(cols)}) VALUES\n" +
            ",\n".join("(" + ",".join(sql_value(r.get(c)) for c in cols) + ")" for r in rows[i:i + chunk]) + ";"
            for i in range(0, len(rows), chunk)]


def to_sql(data: dict, generated_at: str) -> str:
    stamp = {"created_at": generated_at, "updated_at": generated_at}
    parts = ["-- Generated by migration/extract.py. Do not edit by hand.",
             "-- No BEGIN/COMMIT: `wrangler d1 execute --file` rejects explicit transactions.",
             "DELETE FROM issues;", "DELETE FROM import_issues;", "DELETE FROM reservations;",
             "DELETE FROM vessels;", "DELETE FROM berths;", "DELETE FROM app_meta;"]
    parts += inserts("berths", ["id", "name", "length_ft", "is_exclusive", "sort_order", "source"], data["berths"])
    parts += inserts("vessels", ["id", "name", "name_key", "length_ft", "length_status", "length_note", "source",
                                 "created_at", "updated_at"], [{**v, **stamp} for v in data["vessels"]])
    parts += inserts("reservations", ["id", "berth_id", "kind", "vessel_id", "title", "start_date", "end_date",
                                      "notes", "origin", "source_ref", "created_at", "updated_at"],
                     [{**r, **stamp} for r in data["reservations"]])
    parts += inserts("issues", ["id", "type", "reservation_id", "other_reservation_id", "berth_id", "start_date",
                                "end_date", "details", "created_at"], data["issues"])
    parts += inserts("import_issues", ["id", "code", "severity", "sheet", "cell", "message", "reservation_id"],
                     [{**i, "id": n + 1} for n, i in enumerate(data["import_issues"])])
    parts += inserts("app_meta", ["key", "value"], [
        {"key": "import_summary", "value": json.dumps(data["summary"], separators=(",", ":"))},
        {"key": "data_range", "value": json.dumps(data["summary"]["data_range"])},
        {"key": "max_span_days", "value": str(data["summary"]["max_span_days"])}])
    return "\n".join(parts) + "\n"


def main():
    ap = argparse.ArgumentParser(description="Migrate the legacy dock workbook")
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--generated-at", default=None, help="fixed ISO timestamp for reproducible output")
    args = ap.parse_args()
    generated_at = args.generated_at or dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
    data = Extractor(args.input).build(generated_at)
    data["summary"]["source_file"] = args.input.name
    args.out.mkdir(parents=True, exist_ok=True)
    for name in ("berths", "vessels", "reservations", "issues", "import_issues"):
        (args.out / f"{name}.json").write_text(json.dumps(data[name], indent=1))
    (args.out / "import_report.json").write_text(json.dumps(
        {"summary": data["summary"], "import_issues": data["import_issues"]}, indent=1))
    (args.out / "seed.sql").write_text(to_sql(data, generated_at))
    print(json.dumps(data["summary"], indent=2))


if __name__ == "__main__":
    main()
